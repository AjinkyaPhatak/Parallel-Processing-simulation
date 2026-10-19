"""Build report/technical_note.docx (pandoc) and .pdf (Microsoft Word export) from report/technical_note.md,
then report the page count.   Usage:  python scripts/build_report.py   (needs scripts/env.ps1 for pandoc on PATH,
or tools/pandoc/pandoc.exe is used directly)."""
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
REP = ROOT / "report"
MD, DOCX, PDF = REP / "technical_note.md", REP / "technical_note.docx", REP / "technical_note.pdf"
REF = REP / "_reference.docx"
PANDOC = shutil.which("pandoc") or str(ROOT / "tools" / "pandoc" / "pandoc.exe")

# 1. reference document with compact styles
subprocess.run([PANDOC, "-o", str(REF), "--print-default-data-file", "reference.docx"], check=True)
ref = Document(REF)
for sec in ref.sections:
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin = sec.bottom_margin = Cm(1.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
st = ref.styles
for name, size, bold in [("Normal", 9.5, False), ("Body Text", 9.5, False), ("First Paragraph", 9.5, False),
                         ("Compact", 9, False), ("Title", 14, True), ("Subtitle", 10, False), ("Author", 9.5, False),
                         ("Date", 9, False), ("Heading 1", 11, True), ("Heading 2", 10, True),
                         ("Image Caption", 8, False), ("Table Caption", 8, False), ("Caption", 8, False),
                         ("Bibliography", 7.5, False), ("Verbatim Char", 7.5, False)]:
    try:
        s = next(x for x in st if x.name == name)
    except StopIteration:
        continue
    s.font.name = "Consolas" if name == "Verbatim Char" else "Calibri"
    s.font.size = Pt(size)
    s.font.bold = bold
    if name.startswith("Heading") or name == "Title":
        s.font.color.rgb = RGBColor(0x0B, 0x0B, 0x0B)
    if not hasattr(s, "paragraph_format"):   # character style (e.g. inline code): font only
        continue
    pf = s.paragraph_format
    pf.space_before = Pt(6 if name.startswith("Heading") else 0)
    pf.space_after = Pt(2 if name.startswith("Heading") else 3)
    pf.line_spacing = 1.0
    if name in ("Body Text", "First Paragraph"):
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
ref.save(REF)

# 2. markdown -> docx
subprocess.run([PANDOC, str(MD), "-o", str(DOCX), "--reference-doc", str(REF), "--resource-path", str(REP),
                "-f", "markdown+pipe_tables+implicit_figures"], check=True, cwd=REP)

# 3. small post-fix: shrink table text
doc = Document(DOCX)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(7)
doc.save(DOCX)

# 4. docx -> pdf via Word (COM), then count pages
ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
      f"$d = $w.Documents.Open('{DOCX}'); $d.SaveAs([ref]'{PDF}', [ref]17); $d.Close(); $w.Quit()")
r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True)
if r.returncode != 0 or not PDF.exists():
    print("PDF export via Word failed:", r.stderr[-800:])
    sys.exit(1)
from pypdf import PdfReader
n = len(PdfReader(str(PDF)).pages)
print(f"docx: {DOCX}\npdf : {PDF}\npages: {n}")
